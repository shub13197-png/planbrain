"""Manual column mapping: header detection, the Excel realities, and profiles.

This is the fallback the mapping model will sit on top of, so it has to be
genuinely usable on its own. A mapping assistant that is the *only* way to
import fails whenever it is wrong, and nobody can tell when it is wrong without
a manual path to compare against.

Every "Excel reality" below is one that has actually bitten someone: a title
block above the header, a two-row header with a merged group label, a material
number whose leading zero is load-bearing, a unit suffix in the heading, and
trailing whitespace that makes a name match nothing.
"""

from datetime import date

import pytest

from planbrain.mapping import (
    Profile,
    ProfileError,
    apply_mapping,
    discover,
    inspect,
    load,
    normalise,
    read_grid,
    required_fields,
    save,
    sheet_names,
    slug,
    suggest,
)
from planbrain.mapping.io import as_text

pytest.importorskip("openpyxl")


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def _xlsx(tmp_path, name, build):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    build(sheet)
    path = tmp_path / name
    book.save(path)
    return path


def _csv(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


HISTORY_MAP = {"sku_id": "Item Code", "loc_id": "Godown",
               "bucket_date": "Date", "qty": "Qty"}


# --------------------------------------------------------------------------
# the header is not row 1
# --------------------------------------------------------------------------

def test_a_title_block_above_the_header_is_skipped(tmp_path):
    """Company name, report title, date range, blank, THEN the headings. This
    is what an ERP export actually looks like."""
    path = _csv(tmp_path, "report.csv",
                "SHREE LUBRICANTS PVT LTD,,,\n"
                "Monthly Stock Movement,,,\n"
                "Period: April 2026,,,\n"
                ",,,\n"
                "Item Code,Godown,Date,Qty\n"
                "1000,11,2026-04-02,120\n"
                "1000,11,2026-04-03,95\n")
    guess = inspect(read_grid(path))
    assert guess.header_row == 4
    assert guess.columns[:2] == ["item_code", "godown"]
    assert guess.skipped_rows == [0, 1, 2]


def test_the_reasoning_is_returned_not_just_the_answer(tmp_path):
    """A wrong header row shifts every column by one and produces a file that
    imports cleanly and means nothing. The user has to be able to check it."""
    path = _csv(tmp_path, "r.csv",
                "Title,,\nItem Code,Godown,Qty\n1000,11,5\n1001,12,7\n")
    guess = inspect(read_grid(path))
    assert guess.reasons
    assert 0.0 <= guess.confidence <= 1.0
    assert any("numbers" in r or "title" in r for r in guess.reasons)


def test_a_merged_two_row_header_is_joined(tmp_path):
    """'Material' merged across two columns, with 'Item Code' and 'Description'
    beneath it, must not read as one named column and one blank."""
    def build(sheet):
        sheet["A1"] = "Material"
        sheet.merge_cells("A1:B1")
        sheet["A2"] = "Item Code"
        sheet["B2"] = "Description"
        sheet["C2"] = "Qty"
        for i, row in enumerate([("1000", "SN-150", 5), ("1001", "ZDDP", 7)], start=3):
            for j, value in enumerate(row, start=1):
                sheet.cell(row=i, column=j, value=value)

    guess = inspect(read_grid(_xlsx(tmp_path, "merged.xlsx", build)))
    assert guess.raw_columns == ["Material Item Code", "Material Description", "Qty"]
    assert guess.header_rows == [0, 1]


def test_a_merged_cell_does_not_become_blank_columns(tmp_path):
    """openpyxl returns the value only for a merge's top-left cell. Without
    filling it forward, a header merged across three columns loses two names."""
    def build(sheet):
        sheet["A1"] = "Movement"
        sheet.merge_cells("A1:C1")
        for j, name in enumerate(["Date", "Godown", "Qty"], start=1):
            sheet.cell(row=2, column=j, value=name)
        sheet["A3"], sheet["B3"], sheet["C3"] = "2026-04-02", "11", "5"

    grid = read_grid(_xlsx(tmp_path, "m.xlsx", build))
    assert grid[0][:3] == ["Movement", "Movement", "Movement"]


# --------------------------------------------------------------------------
# the value realities
# --------------------------------------------------------------------------

def test_a_leading_zero_survives_reading(tmp_path):
    """Material 007821 must reach the mapping as text. What the canonical field
    then does with it is a separate, visible decision."""
    path = _csv(tmp_path, "z.csv", "Item Code,Qty\n007821,5\n")
    grid = read_grid(path)
    assert as_text(grid[1][0]) == "007821"


def test_a_lost_leading_zero_is_warned_about_not_hidden(tmp_path):
    """`sku_id` is an integer everywhere in this system, so 007821 becomes 7821.
    Usually harmless; a real problem if a catalogue holds both as distinct
    materials. Only the user knows which, so the preview tells them."""
    path = _csv(tmp_path, "z.csv",
                "Item Code,Godown,Date,Qty\n007821,11,2026-04-02,5\n")
    guess = inspect(read_grid(path))
    result = apply_mapping(read_grid(path), HISTORY_MAP, "history",
                           first_data_row=guess.first_data_row,
                           headers=guess.raw_columns)
    assert result.ok
    assert len(result.warnings) == 1
    warning = result.warnings[0]
    assert warning["kind"] == "leading_zero"
    assert warning["example"] == "007821 -> 7821"
    assert "different parts" in warning["message"]


def test_an_excel_float_does_not_gain_a_decimal_point():
    """Excel hands back 7821.0 for a text-formatted cell. str() would make that
    '7821.0', which is a different material number."""
    assert as_text(7821.0) == "7821"
    assert as_text(85.5) == "85.5"


def test_a_unit_suffix_in_the_header_still_matches(tmp_path):
    """'Qty (Kgs)', 'QTY_KG' and 'Quantity' are the same column to a human."""
    for header in ("Qty (Kgs)", "QTY_KG", "Quantity (Kgs)", "Quantity"):
        assert suggest([header, "Item Code", "Godown", "Date"], "history")["qty"] == header


def test_trailing_whitespace_does_not_break_matching():
    assert normalise("Item Code ") == "item_code"
    assert normalise("  Godown ") == "godown"


def test_the_original_header_is_preserved_for_display(tmp_path):
    """Matching is done on a normalised name; the user always sees theirs.
    'Qty (Kgs)' tells a person something 'qty' does not."""
    path = _csv(tmp_path, "u.csv", "Item Code ,Qty (Kgs)\n1000,5\n1001,6\n")
    guess = inspect(read_grid(path))
    assert "Qty (Kgs)" in guess.raw_columns
    assert "qty" in guess.columns


# --------------------------------------------------------------------------
# suggestions are deliberately not clever
# --------------------------------------------------------------------------

def test_a_merged_group_label_does_not_break_a_suggestion():
    """'Material Item Code' is still an item code. The specific part of a
    merged heading is its tail."""
    headers = ["Material Item Code", "Movement Godown", "Movement Posting Date",
               "Qty (Kgs)"]
    assert suggest(headers, "history") == {
        "sku_id": "Material Item Code", "loc_id": "Movement Godown",
        "bucket_date": "Movement Posting Date", "qty": "Qty (Kgs)",
    }


def test_one_column_cannot_supply_two_fields():
    """Without this a file with both 'Date' and 'Posting Date' can map both to
    bucket_date and silently drop one."""
    mapping = suggest(["Date", "Posting Date", "Item Code", "Qty"], "history")
    assert len(set(mapping.values())) == len(mapping)


def test_an_unrecognisable_header_is_left_blank_not_guessed():
    """A plausible wrong guess accepted without reading is worse than a blank
    dropdown, because nobody checks a field that looks filled in."""
    mapping = suggest(["Column1", "Column2", "Column3"], "history")
    assert mapping == {}


def test_required_fields_come_from_the_parsers_not_a_second_list():
    """Two lists would eventually disagree."""
    assert set(required_fields("history")) == {"sku_id", "loc_id", "bucket_date", "qty"}
    assert "safety_stock" not in required_fields("parts")


# --------------------------------------------------------------------------
# preview and the all-or-nothing rule
# --------------------------------------------------------------------------

def test_a_preview_parses_without_writing(tmp_path):
    path = _csv(tmp_path, "h.csv",
                "Item Code,Godown,Date,Qty\n1000,11,2026-04-02,120\n"
                "1000,11,2026-04-03,95\n")
    guess = inspect(read_grid(path))
    result = apply_mapping(read_grid(path), HISTORY_MAP, "history",
                           first_data_row=guess.first_data_row,
                           headers=guess.raw_columns)
    assert result.ok
    assert result.total_rows == 2
    assert result.preview[0] == {"sku_id": 1000, "loc_id": 11,
                                 "bucket_date": date(2026, 4, 2), "qty": 120.0}


def test_an_unmapped_required_field_blocks_the_import(tmp_path):
    path = _csv(tmp_path, "h.csv", "Item Code,Qty\n1000,5\n")
    guess = inspect(read_grid(path))
    result = apply_mapping(read_grid(path), {"sku_id": "Item Code", "qty": "Qty"},
                           "history", first_data_row=guess.first_data_row,
                           headers=guess.raw_columns)
    assert not result.ok
    assert set(result.missing_required) == {"loc_id", "bucket_date"}


def test_mapping_to_a_column_that_is_not_there_is_an_error(tmp_path):
    """A stale profile applied to a changed export. Must say so, not import
    three of four fields."""
    path = _csv(tmp_path, "h.csv", "Item Code,Godown,Date,Qty\n1000,11,2026-04-02,5\n")
    guess = inspect(read_grid(path))
    result = apply_mapping(read_grid(path), {**HISTORY_MAP, "qty": "Quantity"},
                           "history", first_data_row=guess.first_data_row,
                           headers=guess.raw_columns)
    assert not result.ok
    assert "not in this file" in str(result.log.errors[0])


def test_columns_left_unmapped_are_reported(tmp_path):
    """A user should know what is being ignored, not discover it later."""
    path = _csv(tmp_path, "h.csv",
                "Item Code,Godown,Date,Qty,Remarks\n1000,11,2026-04-02,5,ok\n")
    guess = inspect(read_grid(path))
    result = apply_mapping(read_grid(path), HISTORY_MAP, "history",
                           first_data_row=guess.first_data_row,
                           headers=guess.raw_columns)
    assert result.unmapped_sources == ["Remarks"]


def test_one_bad_row_does_not_discard_the_good_ones(tmp_path):
    path = _csv(tmp_path, "h.csv",
                "Item Code,Godown,Date,Qty\n1000,11,2026-04-02,5\n"
                "1001,11,not-a-date,7\n1002,11,2026-04-04,9\n")
    guess = inspect(read_grid(path))
    result = apply_mapping(read_grid(path), HISTORY_MAP, "history",
                           first_data_row=guess.first_data_row,
                           headers=guess.raw_columns)
    assert result.total_rows == 2
    assert len(result.log) == 1
    assert result.log.errors[0].row == 3


def test_a_blank_spacer_row_is_not_an_error(tmp_path):
    path = _csv(tmp_path, "h.csv",
                "Item Code,Godown,Date,Qty\n1000,11,2026-04-02,5\n,,,\n"
                "1001,11,2026-04-03,7\n")
    guess = inspect(read_grid(path))
    result = apply_mapping(read_grid(path), HISTORY_MAP, "history",
                           first_data_row=guess.first_data_row,
                           headers=guess.raw_columns)
    assert result.ok and result.total_rows == 2


# --------------------------------------------------------------------------
# profiles are data
# --------------------------------------------------------------------------

def test_a_hand_written_profile_loads(tmp_path):
    """The whole design constraint: a user can write one in a text editor."""
    path = tmp_path / "mine.yaml"
    path.write_text(
        "name: My depot export\n"
        "table: history\n"
        "header_row: 5\n"
        "columns:\n"
        "  sku_id: Item Code\n"
        "  loc_id: Godown\n"
        "  bucket_date: Date\n"
        "  qty: Quantity\n",
        encoding="utf-8",
    )
    profile = load(path)
    assert profile.name == "My depot export"
    assert profile.columns["sku_id"] == "Item Code"
    assert profile.header_row == 4, "1-based in the file, 0-based internally"


def test_a_profile_round_trips(tmp_path):
    original = Profile(name="Ludhiana monthly", table="history",
                       columns=dict(HISTORY_MAP), header_row=4, sheet="Data")
    path = save(original, tmp_path)
    again = load(path)
    assert again.columns == original.columns
    assert again.header_row == original.header_row
    assert again.sheet == "Data"


def test_a_saved_profile_is_readable_yaml(tmp_path):
    """Not a pickle, not JSON-in-a-string. Someone has to be able to edit it."""
    path = save(Profile(name="X", table="history", columns=dict(HISTORY_MAP)),
                tmp_path)
    text = path.read_text(encoding="utf-8")
    assert "sku_id: Item Code" in text
    assert "table: history" in text


def test_yaml_is_loaded_safely(tmp_path):
    """yaml.load executes arbitrary Python, and a profile is exactly the kind
    of file that gets emailed around."""
    path = tmp_path / "evil.yaml"
    path.write_text(
        "name: evil\ntable: history\n"
        "columns: !!python/object/apply:os.system ['echo pwned']\n",
        encoding="utf-8",
    )
    with pytest.raises(ProfileError):
        load(path)


def test_an_empty_column_map_is_refused(tmp_path):
    """It would load, apply cleanly and import nothing -- the empty-result
    failure this project keeps finding."""
    path = tmp_path / "p.yaml"
    path.write_text("name: x\ntable: history\ncolumns: {}\n", encoding="utf-8")
    with pytest.raises(ProfileError, match="maps nothing"):
        load(path)


def test_a_malformed_profile_names_the_file_and_the_field(tmp_path):
    path = tmp_path / "broken.yaml"
    path.write_text("name: x\ntable: history\n", encoding="utf-8")
    with pytest.raises(ProfileError, match="broken.yaml.*columns"):
        load(path)


def test_a_zero_or_negative_header_row_is_refused(tmp_path):
    """Users count from 1. A 0 means they misunderstood, not row zero."""
    path = tmp_path / "p.yaml"
    path.write_text("name: x\ntable: history\nheader_row: 0\n"
                    "columns:\n  qty: Qty\n", encoding="utf-8")
    with pytest.raises(ProfileError, match="counting from 1"):
        load(path)


def test_discover_reports_a_broken_profile_rather_than_skipping_it(tmp_path):
    """Silently ignoring it means a hand-written profile never appears and the
    user has nothing to debug from."""
    (tmp_path / "good.yaml").write_text(
        "name: good\ntable: history\ncolumns:\n  qty: Qty\n", encoding="utf-8")
    (tmp_path / "bad.yaml").write_text("name: bad\n", encoding="utf-8")
    found, problems = discover(tmp_path)
    assert [p.name for p in found] == ["good"]
    assert len(problems) == 1 and "bad.yaml" in problems[0]


def test_the_shipped_profiles_all_load():
    """A profile we ship that does not parse is worse than one we do not."""
    from planbrain.mapping import BUILTIN_DIR

    found, problems = discover(BUILTIN_DIR)
    assert problems == []
    assert {p.source_hint for p in found} >= {"sap", "tally"}
    for profile in found:
        assert set(profile.columns) >= set(required_fields(profile.table))


def test_slugs_are_safe_filenames():
    assert slug("Ludhiana depot / monthly!") == "ludhiana-depot-monthly"
    assert slug("   ") == "profile"


# --------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------

def test_sheet_names_are_listed_before_anything_else(tmp_path):
    """Picking the wrong sheet is the fastest route to a confidently wrong
    import; workbooks that open on 'Summary' with the data on 'Sheet2' are
    common."""
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Summary"
    book.create_sheet("Stock Ledger")
    path = tmp_path / "two.xlsx"
    book.save(path)
    assert sheet_names(path) == ["Summary", "Stock Ledger"]


def test_a_semicolon_delimited_csv_is_read(tmp_path):
    """European Excel exports use semicolons and nobody mentions it."""
    path = _csv(tmp_path, "eu.csv", "Item Code;Godown;Date;Qty\n1000;11;2026-04-02;5\n")
    grid = read_grid(path)
    assert grid[0] == ["Item Code", "Godown", "Date", "Qty"]


def test_a_missing_file_says_so(tmp_path):
    from planbrain.mapping import UnreadableFile

    with pytest.raises(UnreadableFile):
        read_grid(tmp_path / "nope.csv")
