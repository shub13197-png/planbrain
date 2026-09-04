"""The order list: the rows a planner acts on, and the file they send out.

**This is the gap these tests exist to close.** A full run computes 2001
planned order releases and the application returned none of them -- `plan.run`
reported a forecast summary and a resource utilisation table, and the results
screen rendered utilisation by resource and nothing else. The arithmetic was
right and the answer was discarded at the last step, which is precisely where a
spreadsheet beats it: an Excel planner ends the session with a list they can
sort, print and send to a supplier.

The list is the RELEASE rows, not the receipts. A release says *place this order
on this date*, which is the thing a planner does. Receipts and releases are not
one-to-one -- `_offset` pulls a release back to the previous working bucket and
two receipts can land on the same release date, which is why the demo shows 2070
receipts against 2001 releases -- so a row pairing each release to "its" receipt
date would be inventing a correspondence that does not exist.
"""

import csv
from datetime import date, timedelta
from types import SimpleNamespace

import pytest

from planbrain import netreq, orders
from planbrain.demo import build_demo, populate
from planbrain.demo.generate import Location, Part
from planbrain.netreq.core import ItemPlan, LotSizing


@pytest.fixture
def planned(con):
    """A real plan in the fact table, netted from replayed history."""
    demo = build_demo(seed=7)
    populate(con, demo)
    netreq.run(con, demo)
    return demo


# --------------------------------------------------------------------------
# the identity: what is reported is what was planned
# --------------------------------------------------------------------------

def stored(con, measure, *, scenario_id=0):
    """Count and total one measure straight out of storage.

    **Deliberately raw SQL, and this file is allowlisted in
    `tools/check_fact_access.py` for it.** The accessor is the right way to read
    a measure and it is the wrong instrument here: `read_facts` densifies across
    the bucket spine for the keys it is *given*, so checking `order_list` with
    the same key set it uses would agree with it about any key it failed to ask
    for. A whole location silently missing from the list would pass.

    Counting what is physically in the table is the one check that cannot be
    fooled that way, and it is the same reason the storage-layer tests are
    allowlisted.
    """
    return con.execute(
        f"SELECT count(*), coalesce(sum(qty), 0) FROM fact_supply_demand"
        f" WHERE measure = ? AND scenario_id = ?",
        (measure, scenario_id),
    ).fetchone()


def test_the_order_list_totals_the_planned_releases_it_reports(con, planned):
    """Every planned release reaches the list, exactly once.

    The failure this catches is a list that quietly drops rows or double-counts
    them -- which a reader cannot see, because a list of 2000 orders looks the
    same either way.
    """
    rows = orders.order_list(con, planned)
    count, total = stored(con, "planned_order_release")

    assert count > 0, "the fixture planned nothing; this would pass vacuously"
    assert len(rows) == count
    assert round(sum(o.qty for o in rows), 6) == round(total, 6)


def test_the_order_list_is_ordered_by_the_date_it_must_be_placed(con, planned):
    """A planner works the list top down; the top must be the next thing due."""
    rows = orders.order_list(con, planned)
    assert [o.release_date for o in rows] == sorted(o.release_date for o in rows)


def test_every_order_names_the_item_rather_than_only_its_id(con, planned):
    """`3004` is not something a planner can take to a supplier."""
    rows = orders.order_list(con, planned)
    assert all(o.name and not o.name.isdigit() for o in rows)
    assert {o.action for o in rows} <= {"make", "buy"}


def test_a_raw_part_is_bought_and_a_finished_part_is_made(con, planned):
    """The list splits by who acts on it: purchasing or the plant."""
    level = {p.sku_id: p.level for p in planned.parts}
    rows = orders.order_list(con, planned)
    for o in rows:
        assert o.action == ("buy" if level[o.sku_id] == "raw" else "make")


# --------------------------------------------------------------------------
# textbook fixture: the published Wagner-Whitin lots survive to the list
# --------------------------------------------------------------------------

def test_the_list_carries_the_snyder_shen_lots_unchanged(con):
    """Snyder & Shen Example 3.9: order 210 in period 1, 150 in period 3.

    Asserted here rather than only in the engine because the list is a second
    chance to get the number wrong -- a rounding, a unit, a bucket offset. The
    numbers a planner reads are the numbers the textbook publishes.
    """
    start = date(2026, 7, 1)
    demo = _one_part_demo(start, periods=4)

    plan = ItemPlan(
        sku_id=1, loc_id=1,
        net_req=[90.0, 120.0, 80.0, 70.0],
        planned_order_receipt=[210.0, 0.0, 150.0, 0.0],
        planned_order_release=[210.0, 0.0, 150.0, 0.0],
        projected_on_hand=[120.0, 0.0, 70.0, 0.0],
        exceptions=[],
    )
    netreq.adapters.write_plans(
        con, [plan], scenario_id=0,
        horizon_start=start, horizon_end=start + timedelta(days=3),
        gross_by_sku={1: [90.0, 120.0, 80.0, 70.0]},
    )

    rows = orders.order_list(con, demo)

    assert [(o.release_date, o.qty) for o in rows] == [
        (start, 210.0),
        (start + timedelta(days=2), 150.0),
    ]


# --------------------------------------------------------------------------
# export: the file is the deliverable
# --------------------------------------------------------------------------

def test_the_exported_csv_reads_back_as_the_list_it_was_given(con, planned, tmp_path):
    """A file that loses rows on the way out is worse than no file at all."""
    rows = orders.order_list(con, planned)
    path = tmp_path / "orders.csv"
    written = orders.to_csv(rows, path)

    assert written == len(rows)
    with path.open(newline="", encoding="utf-8") as fh:
        read_back = list(csv.DictReader(fh))

    assert len(read_back) == len(rows)
    assert round(sum(float(r["qty"]) for r in read_back), 6) == round(
        sum(o.qty for o in rows), 6
    )


def test_the_exported_workbook_reads_back_as_the_list_it_was_given(
    con, planned, tmp_path
):
    """Excel is where this list is going; openpyxl already ships for reading."""
    from openpyxl import load_workbook

    rows = orders.order_list(con, planned)
    path = tmp_path / "orders.xlsx"
    written = orders.to_xlsx(rows, path)

    assert written == len(rows)
    sheet = load_workbook(path).active
    header = [c.value for c in sheet[1]]
    body = list(sheet.iter_rows(min_row=2, values_only=True))

    assert len(body) == len(rows)
    qty = header.index("qty")
    assert round(sum(r[qty] for r in body), 6) == round(sum(o.qty for o in rows), 6)


def test_exporting_nothing_refuses_rather_than_writing_an_empty_plan(tmp_path):
    """An empty file must not read as a plan with no orders in it.

    Same shape as the size gate that passed because it measured nothing: a
    planner who opens a header-only sheet concludes there is nothing to order,
    when what actually happened is that no plan was run.
    """
    path = tmp_path / "orders.csv"
    with pytest.raises(ValueError, match="no planned orders"):
        orders.to_csv([], path)
    assert not path.exists()


# --------------------------------------------------------------------------
# the methods the interface calls
# --------------------------------------------------------------------------

def test_the_screen_is_told_how_many_orders_it_is_not_showing(con, planned):
    """A page of 100 out of 2001 must not read as a list of 100."""
    from planbrain.backend.api import Session, plan_orders

    state = Session(con=con, demo=planned)
    r = plan_orders(state, limit=100)

    assert r["shown"] == 100
    assert r["total"] > r["shown"]
    assert len(r["orders"]) == 100
    assert r["totals"]["make"] + r["totals"]["buy"] == r["total"]


def test_the_export_writes_the_whole_list_not_the_page_on_screen(
    con, planned, tmp_path
):
    """The failure this forbids looks complete, which is what makes it bad."""
    from planbrain.backend.api import Session, plan_export, plan_orders

    state = Session(con=con, demo=planned)
    shown = plan_orders(state, limit=10)
    path = tmp_path / "orders.csv"

    written = plan_export(state, str(path))

    assert written["rows"] == shown["total"]
    assert written["rows"] > shown["shown"]


def test_a_filtered_export_writes_the_filter_the_screen_is_showing(
    con, planned, tmp_path
):
    """Filtering to purchasing and exporting must not hand back the plant's work."""
    from planbrain.backend.api import Session, plan_export, plan_orders

    state = Session(con=con, demo=planned)
    buys = plan_orders(state, limit=1, action="buy")["total"]

    written = plan_export(state, str(tmp_path / "buy.csv"), action="buy")
    assert written["rows"] == buys


def test_an_unwritable_format_is_refused_by_name(con, planned, tmp_path):
    """`.pdf` in the save box must say so, not fail somewhere in openpyxl."""
    from planbrain.backend.api import Session, plan_export

    state = Session(con=con, demo=planned)
    with pytest.raises(ValueError, match="name the file"):
        plan_export(state, str(tmp_path / "orders.pdf"))


def test_asking_for_orders_before_a_plan_says_what_to_do(con):
    """The message a user meets first should name the next action."""
    from planbrain.backend.api import Session, plan_orders

    with pytest.raises(ValueError, match="no dataset loaded"):
        plan_orders(Session(con=con, demo=None))


# --------------------------------------------------------------------------

def _one_part_demo(start, *, periods):
    """The reference data the list needs, and nothing else.

    Reference data reaches the engines as an argument because InvenTree core is
    read-only -- see `netreq.run`. A stand-in here keeps the textbook fixture
    about the arithmetic instead of about the demo generator.
    """
    return SimpleNamespace(
        parts=[Part(sku_id=1, name="Widget", level="finished",
                    lead_time_days=0, safety_stock=0.0,
                    lot_policy="wagner_whitin", lot_qty=0.0)],
        locations=[Location(loc_id=1, name="Plant", kind="plant")],
        horizon_start=start,
        horizon_end=start + timedelta(days=periods - 1),
    )
